import streamlit as st
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import pandas as pd
import io

# ==============================================================================
# 1. MATRIZ DE DESCRITORES PEDAGÓGICOS (TEXTOS BASE)
# ==============================================================================
DESCRITORES = {
    "portugues_leitura": {
        "titulo": "Língua Portuguesa (Leitura e Interpretação)",
        1: "Demonstra necessidade de apoio constante para realizar a leitura e interpretação de textos do 5º ano, necessitando de mediação para identificar a ideia principal.",
        2: "Lê e interpreta textos de complexidade mediana com auxílio pontual, estando em processo de consolidação da capacidade de realizar inferências.",
        3: "Lê e interpreta de maneira satisfatória diferentes gêneros textuais, identificando ideias centrais e realizando inferências adequadas à etapa.",
        4: "Demonstra leitura fluente, crítica e autônoma, interpretando textos com profundidade, identificando implícitos e relacionando fontes diversas."
    },
    "portugues_escrita": {
        "titulo": "Língua Portuguesa (Produção Textual)",
        1: "Necessita de apoio frequente para estruturar frases e produzir textos curtos, apresentando desafios na pontuação, ortografia e coesão.",
        2: "Produz textos curtos com autonomia parcial, estando em desenvolvimento no uso correto de pontuação, paragrafação e elementos coesivos.",
        3: "Elabora produções escritas estruturadas de forma satisfatória, demonstrando boa organização de ideias, coerência e uso adequado das normas gramaticais.",
        4: "Produz textos autônomos com excelente coesão, coerência, riqueza vocabular e domínio consistente das convenções da escrita."
    },
    "matematica_operacoes": {
        "titulo": "Matemática (Numeração e Operações)",
        1: "Demonstra dificuldades no sistema de numeração decimal e nas operações básicas do 5º ano, necessitando de material concreto e apoio constante.",
        2: "Compreende o valor posicional e realiza cálculos básicos com mediação pontual, estando em processo de consolidação dos algoritmos.",
        3: "Demonstra domínio satisfatório dos conceitos numéricos e operações, realizando cálculos de forma autônoma e precisa.",
        4: "Possui raciocínio matemático ágil e autônomo, dominando operações complexas e aplicando estratégias diversificadas com segurança."
    },
    "matematica_problemas": {
        "titulo": "Matemática (Resolução de Problemas)",
        1: "Apresenta desafios para interpretar enunciados de problemas matemáticos, necessitando que as etapas sejam guiadas pelo professor.",
        2: "Interpreta problemas simples e aplica estratégias com apoio pontual, trabalhando na seleção da operação adequada para situações complexas.",
        3: "Resolve situações-problema com autonomia, compreendendo enunciados e aplicando estratégias coerentes para chegar à solução.",
        4: "Demonstra excelente capacidade analítica, formulando hipóteses e resolvendo problemas complexos por múltiplos caminhos de forma autônoma."
    },
    "ciencias_natureza": {
        "titulo": "Ciências da Natureza",
        1: "Necessita de mediação frequente para compreender conceitos científicos do 5º ano e relacioná-los com fenômenos do cotidiano.",
        2: "Compreende conceitos fundamentais das Ciências da Natureza com mediação, demonstrando curiosidade em atividades investigativas.",
        3: "Investiga fenômenos e compreende conceitos científicos de maneira satisfatória, articulando teoria e observações cotidianas.",
        4: "Demonstra pensamento científico apurado, formulando hipóteses, investigando fenômenos e tirando conclusões com autonomia consistente."
    },
    "ciencias_humanas": {
        "titulo": "Ciências Humanas (História e Geografia)",
        1: "Compreende relações sociais, temporais e espaciais com apoio frequente, necessitando de mediação para ler mapas e linhas do tempo.",
        2: "Demonstra compreensão em processo sobre as transformações históricas e geográficas, relacionando conteúdos com apoio pontual.",
        3: "Compreende adequadamente as relações entre sociedade, espaço e tempo, analisando transformações históricas e geográficas com clareza.",
        4: "Analisou com autonomia e profundidade a diversidade cultural, histórica e espacial, estabelecendo conexões críticas e fundamentadas."
    },
    "cultura_digital": {
        "titulo": "Cultura Digital e Computação",
        1: "Necessita de mediação e orientação constante para utilizar tecnologias digitais e recursos educacionais no ambiente escolar.",
        2: "Utiliza ferramentas digitais básicas sob supervisão, em processo de desenvolvimento do uso ético e reflexivo das mídias.",
        3: "Utiliza tecnologias digitais de forma responsável e funcional, atendendo às propostas de aprendizagem e pesquisa com autonomia.",
        4: "Demonstra fluência digital e pensamento computacional avançado, utilizando ferramentas com responsabilidade, criatividade e autonomia."
    },
    "autonomia_foco": {
        "titulo": "Organização e Foco nas Atividades",
        1: "Apresenta dispersão frequente e necessita de intervenções constantes para iniciar, manter o foco e concluir as tarefas propostas.",
        2: "Mantém o foco nas atividades por períodos curtos, necessitando de lembretes para gerenciar o tempo e concluir as produções.",
        3: "Inicia e conclui as tarefas no tempo previsto de forma satisfatória, mantendo atenção e organização do seu espaço de trabalho.",
        4: "Demonstra excelente autonomia, gerenciando seu tempo com eficácia, mantendo foco consistente e buscando superar desafios acadêmicos."
    },
    "comportamento_normas": {
        "titulo": "Atitude e Cumprimento de Combinados",
        1: "Apresenta dificuldades recorrentes em seguir os combinados da turma, necessitando de mediação constante para respeitar os momentos de escuta.",
        2: "Segue a rotina e os combinados da turma na maior parte do tempo, necessitando de direcionamento em momentos de transição.",
        3: "Respeita as normas, orientações e combinações da turma de forma satisfatória, promovendo um ambiente favorável ao aprendizado.",
        4: "Exerce papel positivo na turma, cumprindo e incentivando o respeito aos combinados com extrema maturidade e autonomia."
    },
    "socioemocional_cooperacao": {
        "titulo": "Relações Interpessoais e Empatia",
        1: "Enfrenta desafios no trabalho em equipe e na resolução de conflitos, necessitando de intervenção frequente do professor.",
        2: "Relaciona-se bem com os colegas, mas em atividades em grupo necessita de mediação pontual para dividir tarefas e acolher opiniões.",
        3: "Interage de forma respeitosa, coopera em trabalhos em grupo e busca resolver pequenos conflitos por meio do diálogo.",
        4: "Demonstra elevada empatia, excelente espírito cooperativo e liderança positiva, acolhendo a diversidade e mediando divergências."
    }
}

# ==============================================================================
# 2. FUNÇÃO DE GERAÇÃO DO DOCUMENTO WORD (.DOCX)
# ==============================================================================
def gerar_docx_relatorio(dados_aluno, selecoes, texto_qualitativo):
    doc = docx.Document()
    
    # Margens
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        
    # Cores
    NAVY = RGBColor(0x00, 0x33, 0x66)
    HEX_BG_HEADER = "003366"
    HEX_BG_ALT = "F0F4F8"
    
    # Título do Documento
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    run_t = p_title.add_run("RELATÓRIO DE ACOMPANHAMENTO PEDAGÓGICO E FORMATIVO")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(14)
    run_t.bold = True
    run_t.font.color.rgb = NAVY
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    run_s = p_sub.add_run("ENSINO FUNDAMENTAL I — 5º ANO (DIRETRIZES BNCC / DCC-PR)")
    run_s.font.name = "Arial"
    run_s.font.size = Pt(10)
    run_s.bold = True
    
    # Tabela 1: Identificação do Estudante
    t_id = doc.add_table(rows=4, cols=2)
    t_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_id.autofit = False
    
    campos_id = [
        [("Estudante:", dados_aluno['nome']), ("Turma / Turno:", f"{dados_aluno['turma']} - {dados_aluno['turno']}")],
        [("Escola:", dados_aluno['escola']), ("Período:", dados_aluno['periodo'])],
        [("Data de Nascimento:", dados_aluno['data_nasc']), ("Professor(a):", dados_aluno['professor'])],
        [("Possui PEI / Adequação:", dados_aluno['pei']), ("Atendimento AEE:", dados_aluno['aee'])]
    ]
    
    for r_idx, row_data in enumerate(campos_id):
        for c_idx, (label, val) in enumerate(row_data):
            cell = t_id.cell(r_idx, c_idx)
            cell.width = Inches(3.4)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            r_lbl = p.add_run(f"{label} ")
            r_lbl.bold = True
            r_lbl.font.size = Pt(9.5)
            r_val = p.add_run(str(val))
            r_val.font.size = Pt(9.5)
            
            # Sombreamento
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_BG_ALT}"/>')
            cell._tc.get_or_add_tcPr().append(shd)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # Seção 1: Parecer Descritivo do Desenvolvimento
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("1. PARECER DESCRITIVO DO DESENVOLVIMENTO E APRENDIZAGEM")
    r_h1.font.size = Pt(11)
    r_h1.bold = True
    r_h1.font.color.rgb = NAVY
    
    # Gerar parágrafos agrupados por área
    nome_aluno = dados_aluno['nome'].split()[0]
    
    # Parágrafo 1: Desenvolvimento Acadêmico
    p_acad = doc.add_paragraph()
    p_acad.paragraph_format.line_spacing = 1.15
    p_acad.paragraph_format.space_after = Pt(6)
    r_intro = p_acad.add_run(f"No decorrer do período observado, o(a) estudante {nome_aluno} apresentou o seguinte desempenho nas áreas do conhecimento: ")
    
    # Concatenação das áreas
    texto_acad = []
    for chave in ["portugues_leitura", "portugues_escrita", "matematica_operacoes", "matematica_problemas", "ciencias_natureza", "ciencias_humanas", "cultura_digital"]:
        nivel = selecoes.get(chave, 3)
        desc = DESCRITORES[chave][nivel]
        texto_acad.append(f"Em {DESCRITORES[chave]['titulo']}, {desc.lower()}")
        
    p_acad.add_run(" ".join(texto_acad))
    
    # Parágrafo 2: Autonomia, Comportamento e Socioemocional
    p_socio = doc.add_paragraph()
    p_socio.paragraph_format.line_spacing = 1.15
    p_socio.paragraph_format.space_after = Pt(6)
    p_socio.add_run(f"No tocante aos aspectos atitudinais, de autonomia e convivência escolar, observa-se que {nome_aluno}: ")
    
    texto_socio = []
    for chave in ["autonomia_foco", "comportamento_normas", "socioemocional_cooperacao"]:
        nivel = selecoes.get(chave, 3)
        desc = DESCRITORES[chave][nivel]
        texto_socio.append(desc)
        
    p_socio.add_run(" ".join(texto_socio))
    
    # Seção 2: Registros Qualitativos e Encaminhamentos
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("2. REGISTROS QUALITATIVOS E INTERVENÇÕES PEDAGÓGICAS")
    r_h2.font.size = Pt(11)
    r_h2.bold = True
    r_h2.font.color.rgb = NAVY
    
    secoes_q = [
        ("Potencialidades e Destaques:", texto_qualitativo.get("potencialidades", "Demonstra dedicação e bom relacionamento com a comunidade escolar.")),
        ("Necessidades de Apoio e Dificuldades:", texto_qualitativo.get("dificuldades", "Acompanhamento contínuo no fortalecimento de conceitos acadêmicos específicos.")),
        ("Estratégias de Recuperação Contínua:", texto_qualitativo.get("recuperacao", "Retomada de conteúdos por meio de metodologias diversificadas e apoio pedagógico.")),
        ("Encaminhamentos e Ações Futuras:", texto_qualitativo.get("encaminhamentos", "Manutenção da parceria com a família e monitoramento no próximo trimestre."))
    ]
    
    for tit, texto in secoes_q:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_after = Pt(4)
        p_q.paragraph_format.line_spacing = 1.15
        r_t = p_q.add_run(f"• {tit} ")
        r_t.bold = True
        p_q.add_run(texto)
        
    # Assinaturas
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    p_ass = doc.add_paragraph()
    p_ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass.add_run("____________________________________            ____________________________________\n").bold = True
    p_ass.add_run("              Professor(a) Regente                                    Equipe Pedagógica / Direção\n\n")
    p_ass.add_run("Data da Devolutiva à Família: _____ / _____ / _________   Assinatura do Responsável: ______________________")
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ==============================================================================
# 3. INTERFACE INTERATIVA DO STREAMLIT
# ==============================================================================
def main():
    st.set_page_config(page_title="Gerador de Relatórios Pedagógicos - 5º Ano", page_icon="📝", layout="wide")
    
    st.title("📝 Gerador Automático de Relatórios Pedagógicos (5º Ano)")
    st.caption("Ferramenta formativa alinhada à BNCC, DCC-PR e Diretrizes da Educação Inclusiva.")
    
    # Banco de Dados de Alunos Exemplo
    alunos_default = pd.DataFrame([
        {"nome": "Lucas Gabriel Silva", "turma": "5º Ano A", "turno": "Manhã", "escola": "Escola Municipal Educação e Futuro", "data_nasc": "12/04/2015", "professor": "Profª. Ana Maria", "pei": "Não", "aee": "Não"},
        {"nome": "Beatriz Oliveira Santos", "turma": "5º Ano A", "turno": "Manhã", "escola": "Escola Municipal Educação e Futuro", "data_nasc": "05/08/2015", "professor": "Profª. Ana Maria", "pei": "Sim (Transtorno do Espectro Autista)", "aee": "Sim"},
        {"nome": "Carlos Eduardo Lima", "turma": "5º Ano B", "turno": "Tarde", "escola": "Escola Municipal Educação e Futuro", "data_nasc": "22/11/2014", "professor": "Prof. Roberto Carlos", "pei": "Não", "aee": "Não"}
    ])
    
    st.sidebar.header("⚙️ 1. Seleção do Estudante")
    
    # Upload de Lista de Alunos (Opcional)
    uploaded_file = st.sidebar.file_uploader("Carregar lista de alunos (Excel/CSV)", type=["csv", "xlsx"])
    if uploaded_file is not None:
        try:
            if uploaded_file.name.endswith(".csv"):
                df_alunos = pd.read_csv(uploaded_file)
            else:
                df_alunos = pd.read_excel(uploaded_file)
            st.sidebar.success("Lista de alunos carregada com sucesso!")
        except Exception as e:
            st.sidebar.error(f"Erro ao ler arquivo: {e}")
            df_alunos = alunos_default
    else:
        df_alunos = alunos_default
        
    # Seleção do Aluno no Menu Suspenso
    lista_nomes = df_alunos["nome"].tolist()
    nome_selecionado = st.sidebar.selectbox("Escolha o Estudante:", lista_nomes)
    
    # Obter Dados do Aluno Selecionado
    aluno_info = df_alunos[df_alunos["nome"] == nome_selecionado].iloc[0].to_dict()
    periodo_obs = st.sidebar.radio("Período de Observação:", ["1º Trimestre", "2º Trimestre", "3º Trimestre"])
    aluno_info["periodo"] = periodo_obs
    
    st.subheader(f"👤 Dados do Estudante: {aluno_info['nome']}")
    col1, col2, col3 = st.columns(3)
    with col1:
        st.write(f"**Turma/Turno:** {aluno_info['turma']} ({aluno_info['turno']})")
        st.write(f"**Escola:** {aluno_info['escola']}")
    with col2:
        st.write(f"**Data de Nasc.:** {aluno_info['data_nasc']}")
        st.write(f"**Professor(a):** {aluno_info['professor']}")
    with col3:
        st.write(f"**Possui PEI:** {aluno_info['pei']}")
        st.write(f"**Atendimento AEE:** {aluno_info['aee']}")
        
    st.markdown("---")
    st.header("📊 2. Avaliação dos Indicadores de Desempenho (Escala 1 a 4)")
    st.info("1: Apoio Frequente | 2: Consolidação Parcial | 3: Autonomia Adequada | 4: Autonomia Consistente")
    
    selecoes_usuario = {}
    
    # Criar abas para organizar as seções
    tab1, tab2, tab3 = st.tabs(["📚 Áreas do Conhecimento", "🌱 Autonomia e Comportamento", "✍️ Registros Qualitativos"])
    
    with tab1:
        st.subheader("Desempenho Curricular (BNCC / DCC-PR)")
        keys_acad = ["portugues_leitura", "portugues_escrita", "matematica_operacoes", "matematica_problemas", "ciencias_natureza", "ciencias_humanas", "cultura_digital"]
        for key in keys_acad:
            info = DESCRITORES[key]
            val = st.slider(f"**{info['titulo']}**", min_value=1, max_value=4, value=3, help=f"1: {info[1]}\n4: {info[4]}")
            st.caption(f"*Nível {val}:* {info[val]}")
            selecoes_usuario[key] = val
            st.markdown("<hr style='margin:5px 0;'>", unsafe_allow_html=True)
            
    with tab2:
        st.subheader("Autonomia, Comportamento e Socioemocional")
        keys_socio = ["autonomia_foco", "comportamento_normas", "socioemocional_cooperacao"]
        for key in keys_socio:
            info = DESCRITORES[key]
            val = st.slider(f"**{info['titulo']}**", min_value=1, max_value=4, value=3)
            st.caption(f"*Nível {val}:* {info[val]}")
            selecoes_usuario[key] = val
            st.markdown("<hr style='margin:5px 0;'>", unsafe_allow_html=True)
            
    with tab3:
        st.subheader("Anotações Complementares do Professor")
        potencialidades = st.text_area("Potencialidades e Destaques do Estudante:", value="Demonstra dedicação, boa participação nas aulas e excelente convivência com os colegas.")
        dificuldades = st.text_area("Necessidades de Apoio e Dificuldades Observadas:", value="Acompanhamento contínuo na consolidação da escrita e atenção aos detalhes das operações matemáticas.")
        recuperacao = st.text_area("Estratégias de Recuperação Contínua Realizadas:", value="Atividades de retomada com grupos de apoio e mediação individualizada nas tarefas escritas.")
        encaminhamentos = st.text_area("Encaminhamentos para o Próximo Período:", value="Manutenção do acompanhamento pedagógico e estreitamento da parceria com a família.")
        
        texto_qualitativo = {
            "potencialidades": potencialidades,
            "dificuldades": dificuldades,
            "recuperacao": recuperacao,
            "encaminhamentos": encaminhamentos
        }
        
    st.markdown("---")
    st.header("📄 3. Gerar e Baixar Relatório Formatado")
    
    if st.button("🚀 Processar e Gerar Documento Word (.docx)", type="primary"):
        docx_buffer = gerar_docx_relatorio(aluno_info, selecoes_usuario, texto_qualitativo)
        
        nome_arquivo = f"Relatorio_{aluno_info['nome'].replace(' ', '_')}_{aluno_info['periodo'].replace(' ', '_')}.docx"
        
        st.success(f"✅ Relatório de **{aluno_info['nome']}** gerado com sucesso!")
        
        st.download_button(
            label="📥 Clique para Baixar o Relatório (.docx)",
            data=docx_buffer,
            file_name=nome_arquivo,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

if __name__ == "__main__":
    main()
